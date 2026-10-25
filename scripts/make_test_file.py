"""Generate tests/test_data.xlsx in the exact template format, with every tricky case.

Starts from data/template.xlsx (the official template, including its grey example
row 3, which the parser must skip) and writes data from row 4.

Case map (expected outcome):
  APP-001  Bengaluru, strong, pin 560041                     -> Approved (wins pin over APP-013)
  APP-002  Bengaluru, pin 560004, 0.4 km from ST-03          -> Waitlisted, high cannibalisation flag
  APP-003  Mysuru, rented 450 sq ft                           -> Waitlisted (5 eligible, 2 slots)
  APP-004  Mysuru, owned, investment typed "₹38 lakh"         -> Approved (beats APP-009)
  APP-005  City Mysuru but pin 560045 (Bengaluru)             -> On hold, CITY_PIN_MISMATCH
  APP-006  Mysuru, owned mall unit                            -> Approved (withdraw it in chat)
  APP-007  Rented, size blank, notes blank, 0 experience      -> On hold, SPACE_DETAILS_MISSING
  APP-008  Notes "still arranging funds, bank loan pending"   -> On hold, FUNDING_UNCONFIRMED
  APP-009  X: rented 400 sq ft Tier 2, ₹35 lakh               -> Waitlisted, higher investment score than APP-010
  APP-010  Y: rented 400 sq ft Tier 2, ₹26 lakh               -> Waitlisted
  APP-011  Belagavi pin 590006, no territory                  -> Not approved, OUTSIDE_TERRITORY
  APP-012  Upper-floor rented space                           -> On hold, FLOOR_UNSUITABLE ("found a ground-floor space" in chat)
  APP-013  Same new pin 560041 as APP-001                     -> Waitlisted for the pin
  APP-014  Pin 570009 (existing store ST-06)                  -> Not approved, PIN_HAS_STORE
  APP-015  90% of required investment                          -> On hold, FUNDING_GAP ("can now invest ₹30 lakh" in chat)
  APP-016  Davanagere, ₹30 lakh                                -> Approved (exact score tie with APP-017, wins on ratio)
  APP-017  Davanagere, ₹28 lakh                                -> Waitlisted (tie-break level 2)
  APP-018  Tumakuru, no space, target launch 1 month          -> On hold, NO_SPACE + timeline gap
  APP-019  Hubballi, rented ground 450, size typed "450 sq ft" -> Approved (displaced when APP-012 fixes floor)
  APP-020  Hubballi-Dharwad, 900 sq ft owned                   -> Approved, SPACE_OVERSIZE warning
  APP-021  Duplicate of APP-003 (same name + pin)              -> Not approved, DUPLICATE_APPLICANT
  APP-022  "Not yet identified" but size 450, experience "NA"  -> On hold, NO_SPACE + SPACE_NONE_BUT_DETAILS
  APP-023  Tumakuru, 72% of required                           -> Not approved, FUNDING_SHORT
  APP-024  Tumakuru, 0 experience, Hindi                       -> Approved, Foundation track, shadowing in Bengaluru
  APP-025  Mangaluru, ₹34 lakh, 350 sq ft                      -> Approved (same money as APP-026, smaller space)
  APP-026  Mangaluru, ₹34 lakh, 600 sq ft                      -> Waitlisted (lower investment score)
  APP-027  Bengaluru, mall unit                                -> Approved
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

APPLICATIONS = [
    # id, name, city, state, pin, investment, exp, business, space, size, floor, launch, language, notes
    ("APP-001", "Anitha Rao", "Bengaluru", "Karnataka", 560041, 48, 8, "Runs an optical shop in Jayanagar for 8 years", "Owned", 500, "Ground", 3, "Kannada", "Own building on main road, funds ready from savings"),
    ("APP-002", "Mohammed Irfan", "Bengaluru", "Karnataka", 560004, 40, 3, "Ran a mobile accessories shop for 3 years", "Rented", 400, "Ground", 4, "Hindi", "Shop is in Basavanagudi near the market"),
    ("APP-003", "Suresh Gowda", "Mysuru", "Karnataka", 570002, 32, 4, "Runs a footwear store", "Rented", 450, "Ground", 3, "Kannada", "Wants to open before Diwali"),
    ("APP-004", "Lakshmi Narayan", "Mysuru", "Karnataka", 570008, "₹38 lakh", 6, "Family runs a saree showroom", "Owned", 400, "Ground", 4, "Kannada", "Ground floor of our own building"),
    ("APP-005", "Deepak Shetty", "Mysuru", "Karnataka", 560045, 35, 5, "Runs a stationery store", "Rented", 400, "Ground", 3, "English", ""),
    ("APP-006", "Priya Menon", "Mysuru", "Karnataka", 570011, 40, 6, "Manages a supermarket", "Owned", 420, "Mall unit", 2, "English", "Unit in a new mall, self-funded"),
    ("APP-007", "Ramesh Kulkarni", "Mangaluru", "Karnataka", 575003, 30, 0, "Retired bank employee", "Rented", None, "Ground", 3, "Kannada", None),
    ("APP-008", "Fathima Begum", "Mangaluru", "Karnataka", 575002, 30, 5, "Runs a garments shop", "Rented", 400, "Ground", 3, "Hindi", "still arranging funds, bank loan pending"),
    ("APP-009", "Kiran Hegde", "Mysuru", "Karnataka", 570010, 35, 2, "Runs a bakery", "Rented", 400, "Ground", 3, "Kannada", ""),
    ("APP-010", "Sunil Patil", "Mysuru", "Karnataka", 570020, 26, 2, "Runs a bakery", "Rented", 400, "Ground", 3, "Kannada", ""),
    ("APP-011", "Vinay Desai", "Belagavi", "Karnataka", 590006, 40, 5, "Runs an electronics store", "Owned", 450, "Ground", 3, "Kannada", ""),
    ("APP-012", "Arjun Naik", "Hubballi-Dharwad", "Karnataka", 580031, 40, 7, "Runs a mobile phone shop", "Rented", 500, "Upper", 4, "Kannada", "Space is on the first floor above a bank"),
    ("APP-013", "Kavya Reddy", "Bengaluru", "Karnataka", 560041, 42, 5, "Runs an electronics store", "Rented", 450, "Ground", 3, "English", ""),
    ("APP-014", "Manjunath S", "Mysuru", "Karnataka", 570009, 40, 4, "Runs a gift shop", "Owned", 500, "Ground", 3, "Kannada", ""),
    ("APP-015", "Ganesh Bhat", "Hubballi-Dharwad", "Karnataka", 580021, 22.5, 6, "Runs a tea stall and snacks counter", "Rented", 400, "Ground", 3, "Kannada", ""),
    ("APP-016", "Shilpa Angadi", "Davanagere", "Karnataka", 577002, 30, 3, "Runs a stationery shop", "Owned", 400, "Ground", 3, "Kannada", ""),
    ("APP-017", "Basavaraj H", "Davanagere", "Karnataka", 577004, 28, 3, "Runs a garments shop", "Owned", 400, "Ground", 3, "Kannada", ""),
    ("APP-018", "Naveen Kumar", "Tumakuru", "Karnataka", 572101, 22, 0, "Software engineer", "Not yet identified", None, "Not known", 1, "English", "Want to open next month, still looking for a shop"),
    ("APP-019", "Rekha Joshi", "Hubballi-Dharwad", "Karnataka", 580032, 32, 4, "Runs a medical store", "Rented", "450 sq ft", "Ground", 3, "Kannada", ""),
    ("APP-020", "Imran Khan", "Hubballi-Dharwad", "Karnataka", 580001, 45, 2, "Owns a furniture showroom", "Owned", 900, "Ground", 4, "Hindi", "Large showroom space, will use part of it if needed"),
    ("APP-021", "Suresh Gowda", "Mysuru", "Karnataka", 570002, 32, 4, "Footwear store", "Rented", 450, "Ground", None, "Kannada", None),
    ("APP-022", "Pooja Shenoy", "Mangaluru", "Karnataka", 575004, 28, "NA", "Runs a boutique", "Not yet identified", 450, "Ground", 4, "English", "Have seen a shop, not finalised yet"),
    ("APP-023", "Raghavendra T", "Tumakuru", "Karnataka", 572102, 13, 2, "Runs a kirana store", "Owned", 400, "Ground", 3, "Kannada", ""),
    ("APP-024", "Sunita Verma", "Tumakuru", "Karnataka", 572103, 25, 0, "Worked as a school teacher", "Rented", 350, "Ground", 4, "Hindi", "Funds ready"),
    ("APP-025", "Harish Poojary", "Mangaluru", "Karnataka", 575006, 34, 3, "Runs a mobile recharge shop", "Rented", 350, "Ground", 3, "Kannada", ""),
    ("APP-026", "Nikhil Kamath", "Mangaluru", "Karnataka", 575003, 34, 3, "Runs a sports goods shop", "Rented", 600, "Ground", 3, "English", ""),
    ("APP-027", "Divya Iyer", "Bengaluru", "Karnataka", "560 070", 45, 5, "Runs a watch showroom", "Owned", 450, "Mall unit", 3, "english", ""),
]

TERRITORIES = [
    ("TER-01", "Bengaluru", "Karnataka", "560041, 560070, 560004, 560011, 560095, 560034", "Tier 1", 2, "High"),
    ("TER-02", "Mysuru", "Karnataka", "570002, 570004, 570008, 570009, 570010, 570011, 570020, 570023", "Tier 2", 2, "High"),
    ("TER-03", "Mangaluru", "Karnataka", "575002, 575003, 575004, 575006, 575001", "Tier 2", 1, "Medium"),
    ("TER-04", "Hubballi-Dharwad", "Karnataka", "580001, 580020, 580021, 580029, 580031, 580032", "Tier 2", 2, "Medium"),
    ("TER-05", "Tumakuru", "Karnataka", "572101, 572102, 572103", "Tier 3", 1, "Low"),
    ("TER-06", "Davanagere", "Karnataka", "577001, 577002, 577004", "Tier 3", 1, "Medium"),
]


def build(out: Path = ROOT / "tests" / "test_data.xlsx") -> Path:
    wb = openpyxl.load_workbook(ROOT / "data" / "template.xlsx")
    ws = wb["Applications"]
    for i, row in enumerate(APPLICATIONS):
        for j, v in enumerate(row):
            ws.cell(row=4 + i, column=1 + j, value=v if v != "" else None)
    ws = wb["Territories"]
    for i, row in enumerate(TERRITORIES):
        for j, v in enumerate(row):
            ws.cell(row=4 + i, column=1 + j, value=v)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(f"Wrote {build()}")
