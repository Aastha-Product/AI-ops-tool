"""Build data/sample_test_data.xlsx: the test workbook plus an "Expected results" sheet.

This is the file the web app serves from "Download sample file". It copies
tests/test_data.xlsx (template format, so it uploads as-is) and adds a sheet
listing what each applicant tests, the decision the engine gives, and chat
messages to try. Re-run after changing the test data or the rules:

    python scripts/make_sample_download.py
"""
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "sample_test_data.xlsx"
sys.path.insert(0, str(ROOT))
from engine import parser, pipeline

wb = openpyxl.load_workbook(ROOT / "tests" / "test_data.xlsx")
p = parser.read_workbook(ROOT / "tests" / "test_data.xlsx")
r = pipeline.run_all(p["applications"], p["territories"], p["stores"])
WHAT = {
 "APP-001": "Clean, strong applicant; shares pin 560041 with APP-013", "APP-002": "Only 0.4 km from existing store ST-03 (cannibalisation)",
 "APP-003": "One of 5 eligible applicants for 2 slots in TER-02", "APP-004": "Investment typed as text: '₹38 lakh'",
 "APP-005": "City says Mysuru but pin 560045 is Bengaluru", "APP-006": "Approved; try 'APP-006 has withdrawn' in chat",
 "APP-007": "Rented but space size blank; notes blank; 0 experience", "APP-008": "Notes: 'still arranging funds, bank loan pending'",
 "APP-009": "Same store as APP-010 but more money (₹35 vs ₹26 lakh)", "APP-010": "Same store as APP-009 but less money",
 "APP-011": "Belagavi pin 590006 is in no territory", "APP-012": "Upper-floor space; try 'APP-012 has now found a ground-floor space'",
 "APP-013": "Same new pin 560041 as APP-001", "APP-014": "Pin 570009 already has store ST-06",
 "APP-015": "Has 90% of the money needed; try 'APP-015 can now invest ₹30 lakh'", "APP-016": "Exact score tie with APP-017 (more money)",
 "APP-017": "Exact score tie with APP-016 (less money)", "APP-018": "No space yet and wants to open in 1 month",
 "APP-019": "Approved now; gets displaced when APP-012 fixes its floor", "APP-020": "900 sq ft space, larger than the standard format",
 "APP-021": "Duplicate of APP-003 (same name and pin)", "APP-022": "'Not yet identified' space but a size of 450 is given; experience 'NA'",
 "APP-023": "Has only 72% of the money needed", "APP-024": "No experience, Hindi, Tumakuru (no store in city)",
 "APP-025": "Same money as APP-026 (₹34 lakh), smaller 350 sq ft store", "APP-026": "Same money as APP-025, bigger 600 sq ft store",
 "APP-027": "Bengaluru mall unit; pin typed as '560 070'",
}
ws = wb.create_sheet("Expected results")
ws.append(["Sample test data – what each applicant tests and the result the app should give"])
ws["A1"].font = Font(bold=True, size=12)
ws.append([])
ws.append(["Application ID", "Applicant name", "What it tests", "Expected decision", "Score", "Rank", "Why (app's one-line reason)"])
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
fills = {"Approved": "D1FAE5", "Waitlisted": "FEF3C7", "On hold": "DBEAFE", "Not approved": "FEE2E2"}
for aid, rec in r["records"].items():
    ws.append([aid, rec["name"], WHAT.get(aid, ""), rec["decision"], rec["score"], rec["rank_text"] or "", rec["headline"]])
    ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=fills[rec["decision"]])
c = r["counts"]
ws.append([]); ws.append(["Totals", f"{c['Approved']} Approved, {c['Waitlisted']} Waitlisted, {c['On hold']} On hold, {c['Not approved']} Not approved"])
ws.append([]); ws.append(["Chat tests to try", "", "What should happen"])
for col in (1, 3): ws.cell(ws.max_row, col).font = Font(bold=True)
for q, a in [
 ("Why is APP-007 on hold?", "Explains the missing space size and what to send."),
 ("Why did APP-004 beat APP-009?", "Names the deciding factor (retail experience 15 vs 7)."),
 ("Who is waitlisted in TER-02?", "Lists APP-003, APP-009 and APP-010 with their ranks."),
 ("APP-012 has now found a ground-floor space", "APP-012 On hold → Approved; APP-019 Approved → Waitlisted."),
 ("APP-015 can now invest ₹30 lakh", "APP-015 On hold → Waitlisted (ranks 3rd in TER-04)."),
 ("APP-006 has withdrawn", "APP-006 → Not approved; APP-009 moves up to Approved."),
 ("APP-008's bank loan has been sanctioned", "APP-008 On hold → Waitlisted."),
 ("APP-018 has found a space", "Stays On hold and asks for the space size and floor."),
 ("What if APP-018 found a ground floor shop of 450 sq ft?", "Shows the result without saving it."),
 ("APP-099 found space", "Replies: There is no APP-099 in the uploaded file."),
 ("Reset to uploaded file (button)", "All decisions return to the uploaded file."),
]:
    ws.append([q, "", a])
for col, w in zip("ABCDEFG", [44, 20, 58, 16, 8, 18, 90]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=4):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT)
print("saved", OUT, c)
