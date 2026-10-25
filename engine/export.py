"""Results workbook (openpyxl only, so the web deployment needs no pandas)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

FILLS = {
    "Approved": PatternFill("solid", fgColor="D1FAE5"),
    "Waitlisted": PatternFill("solid", fgColor="FEF3C7"),
    "On hold": PatternFill("solid", fgColor="DBEAFE"),
    "Not approved": PatternFill("solid", fgColor="FEE2E2"),
}


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


def export_xlsx(result: dict, change_log: list[dict] | None = None) -> bytes:
    recs = list(result["records"].values())
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Decisions",
                ["ID", "Name", "City", "Pin", "Territory", "Required ₹ lakh", "Offered ₹ lakh", "Ratio", "Score", "Rank",
                 "Decision", "Why (one line)", "All reasons", "Next steps"],
                [[r["app_id"], r["name"], r["city"], r["pin"], r["territory_id"], r["required_lakh"], r["offered_lakh"],
                  r["ratio"], r["score"], r["rank_text"], r["decision"], r["headline"],
                  " | ".join(x["sentence"] for x in r["reasons"]), " | ".join(r["next_steps"])] for r in recs],
                [10, 18, 14, 9, 10, 10, 10, 7, 7, 18, 13, 60, 80, 80])
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        for c in row:
            if c.value in FILLS:
                c.fill = FILLS[c.value]

    t = result["territories"]
    _sheet(wb, "Territories",
           ["Territory", "City", "Tier", "Priority", "Slots", "Applicants", "Approved", "Waitlisted", "On hold",
            "Not approved", "Pins approved", "Existing-store pins", "Slots left"],
           [[x["territory_id"], x["city"], x["tier"], x["priority"], x["slots"], x["applicants"], x["approved"],
             x["waitlisted"], x["on_hold"], x["not_approved"], x["pins_used"], x["existing_store_pins"],
             x["open_slots_left"]] for x in t])
    _sheet(wb, "Score breakdown", ["ID", "Factor", "Input", "Points", "Max", "Why"],
           [[r["app_id"], b["label"], b["input"], b["points"], b["max"], b["why"]]
            for r in recs if r["score_breakdown"] for b in r["score_breakdown"]], [10, 26, 40, 8, 6, 80])
    _sheet(wb, "Onboarding checklist", ["ID", "Step", "Task", "Timing", "Dates", "Done when", "Detail"],
           [[r["app_id"], s["step"], s["name"], s["timing"], s["dates"], s["done_when"], s["detail"]]
            for r in recs if r["checklist"] for s in r["checklist"]["steps"]], [10, 6, 50, 14, 28, 30, 60])
    _sheet(wb, "Training plans", ["ID", "Track", "Language", "Duration", "Window", "Shadowing", "Modules"],
           [[r["app_id"], r["training"]["track"], r["training"]["language"], r["training"]["duration_text"],
             r["training"]["window"], r["training"]["shadow_store"]["text"] if r["training"]["shadow_store"] else "",
             "; ".join(r["training"]["modules"])] for r in recs if r["training"]], [10, 12, 10, 26, 28, 60, 100])
    _sheet(wb, "Change log", ["Time", "Message", "Applied", "Decision changes"],
           [[e.get("time"), e.get("message"), "; ".join(e.get("updates", [])), "; ".join(e.get("changes", []))]
            for e in (change_log or [])] or [["", "No chat updates", "", ""]], [10, 60, 80, 60])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
