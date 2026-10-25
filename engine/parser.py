"""Excel reading and normalisation.

Reads the judges' workbook as-is: finds the Applications / Territories / Existing
stores sheets, locates the header row, maps headers loosely, normalises every
cell and never crashes on a bad row (it records a warning and moves on).
"""
from __future__ import annotations

import difflib
import io
import math
import re
from typing import Any

import openpyxl

from . import refdata

MISSING_TOKENS = {"", "na", "n/a", "n.a.", "-", "--", "nil", "none", "null", "nan", "not applicable", "?"}

# Normalised header text -> field key. Normalisation: lowercase, ₹ · () / etc. -> space.
APP_HEADERS = {
    "app_id": ["application id", "app id", "application no", "application number", "application", "id"],
    "name": ["applicant name", "name", "applicant", "partner name"],
    "city_state_pin": ["city state pin code", "city state pin", "city state pincode", "location"],
    "city_state": ["city state"],
    "city": ["city"],
    "state": ["state"],
    "pin": ["pin code", "pincode", "pin", "postal code", "zip"],
    "investment": ["investment capacity lakh", "investment capacity rs lakh", "investment capacity", "investment lakh", "investment"],
    "experience": ["retail experience years", "retail experience", "experience years", "experience"],
    "business": ["current past business", "current business", "past business", "business", "occupation"],
    "space": ["space available", "space", "space status"],
    "size": ["space size sq ft", "space size sqft", "space size", "size sq ft", "size", "area sq ft", "area"],
    "floor": ["floor"],
    "launch": ["target launch months", "target launch", "launch months", "launch"],
    "language": ["preferred language", "training language", "language"],
    "notes": ["notes from applicant", "notes", "applicant notes", "remarks", "comments"],
}
TER_HEADERS = {
    "territory_id": ["territory id", "territory", "id"],
    "city_state": ["city state"],
    "city": ["city"],
    "state": ["state"],
    "pins": ["pin codes covered", "pin codes", "pincodes covered", "pins covered", "pins", "pin code"],
    "tier": ["city tier", "tier"],
    "slots": ["open slots", "slots", "open slot"],
    "priority": ["priority"],
}
STORE_HEADERS = {
    "store_id": ["store id", "store"],
    "city": ["city"],
    "pin": ["pin code", "pin", "pincode"],
    "opened": ["opened year", "opened", "year opened", "year"],
}
# Keyword fallbacks, checked in order, for headers no alias matched.
APP_KEYWORDS = [
    ("city_state_pin", ("city", "pin")), ("investment", ("invest",)), ("experience", ("experience",)),
    ("business", ("business",)), ("size", ("size",)), ("space", ("space",)), ("floor", ("floor",)),
    ("launch", ("launch",)), ("language", ("language",)), ("notes", ("note",)), ("pin", ("pin",)),
    ("app_id", ("application",)), ("name", ("name",)),
]
TER_KEYWORDS = [("pins", ("pin",)), ("slots", ("slot",)), ("tier", ("tier",)), ("priority", ("priorit",)),
                ("territory_id", ("territor",))]
STORE_KEYWORDS = [("store_id", ("store",)), ("pin", ("pin",)), ("opened", ("open",))]

CITY_ALIASES = {
    "bengaluru": "Bengaluru", "bangalore": "Bengaluru", "blr": "Bengaluru", "bengaluru urban": "Bengaluru",
    "mysuru": "Mysuru", "mysore": "Mysuru",
    "mangaluru": "Mangaluru", "mangalore": "Mangaluru",
    "hubballi dharwad": "Hubballi-Dharwad", "hubli dharwad": "Hubballi-Dharwad", "hubballi": "Hubballi-Dharwad",
    "hubli": "Hubballi-Dharwad", "dharwad": "Hubballi-Dharwad", "hubballi dharwar": "Hubballi-Dharwad",
    "belagavi": "Belagavi", "belgaum": "Belagavi",
    "tumakuru": "Tumakuru", "tumkur": "Tumakuru",
    "davanagere": "Davanagere", "davangere": "Davanagere",
}

# The template ships one grey example row in each sheet. It is not data.
TEMPLATE_EXAMPLES = {
    "applications": {"app_id": "APP-001", "name": "Ravi Kumar", "business": "Runs a footwear store",
                     "notes": "Wants to open before Diwali"},
    "territories": {"territory_id": "TER-01", "pins_text": "570001, 570002, 570008"},
}


# ---------------------------------------------------------------- cell helpers
def norm_header(text: Any) -> str:
    s = str(text or "").lower().replace("₹", " ").replace("·", " ").replace("•", " ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def is_missing(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    return isinstance(v, str) and v.strip().lower() in MISSING_TOKENS


def clean_text(v: Any) -> str | None:
    if is_missing(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def parse_number(v: Any) -> tuple[float | None, str]:
    """Return (value, status) where status is 'ok', 'missing' or 'invalid'.

    Strips lakh / L / ₹ / sq ft / years / commas. 'crore' multiplies by 100
    (the column is in lakh)."""
    if is_missing(v):
        return None, "missing"
    if isinstance(v, bool):
        return None, "invalid"
    if isinstance(v, (int, float)):
        return float(v), "ok"
    s = str(v).strip().lower().replace(",", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None, "invalid"
    val = float(m.group())
    if "crore" in s or re.search(r"\bcr\b", s):
        val *= 100
    return val, "ok"


def parse_pin(v: Any) -> str | None:
    if is_missing(v):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and not v.is_integer():
            return str(v)
        return str(int(v))
    s = re.sub(r"\s+", "", str(v)).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def canon_city(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    key = norm_header(s)
    if key in CITY_ALIASES:
        return CITY_ALIASES[key]
    close = difflib.get_close_matches(key, list(CITY_ALIASES), n=1, cutoff=0.85)
    return CITY_ALIASES[close[0]] if close else s


def canon_state(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    return "Karnataka" if norm_header(s) in ("karnataka", "ka", "karnatak") else s


def canon_space(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    k = s.lower()
    if any(w in k for w in ("not yet", "not identified", "no space", "none", "looking", "search", "not found")) or k in ("no", "not"):
        return "Not yet identified"
    if "rent" in k or "lease" in k:
        return "Rented"
    if "own" in k:
        return "Owned"
    return s  # unrecognised; validation flags it


def canon_floor(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    k = s.lower()
    if "mall" in k:
        return "Mall unit"
    if "not known" in k or "unknown" in k or "don" in k:
        return "Not known"
    if "ground" in k or k in ("gf", "g/f", "g", "0"):
        return "Ground"
    if any(w in k for w in ("upper", "first", "1st", "second", "2nd", "third", "3rd", "top", "above")):
        return "Upper"
    return s


def canon_language(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    k = s.lower()
    for lang in refdata.config()["languages"]:
        if lang.lower()[:3] in k:
            return lang
    return s


def canon_tier(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    m = re.search(r"[123]", s)
    return f"Tier {m.group()}" if m else s


def canon_priority(v: Any) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    k = s.lower()
    for p, keys in (("High", ("high", "hi")), ("Medium", ("med", "mid")), ("Low", ("low", "lo"))):
        if any(k.startswith(x) for x in keys):
            return p
    return s


def split_location(text: str | None) -> tuple[str | None, str | None, str | None]:
    """'Mysuru · Karnataka · 570001' -> (city, state, pin). Order-insensitive."""
    if not text:
        return None, None, None
    pin = None
    m = re.search(r"\d{3}\s?\d{3}|\d{4,7}", text)
    if m:
        pin = re.sub(r"\s", "", m.group())
        text = text[: m.start()] + " " + text[m.end():]
    parts = [p.strip() for p in re.split(r"[·•|;,/]", text) if p.strip()]
    city = state = None
    for p in parts:
        c = canon_city(p)
        if c in refdata.cities() and city is None:
            city = c
        elif canon_state(p) == "Karnataka" and state is None:
            state = "Karnataka"
    if city is None:  # try free text search
        low = norm_header(text)
        for alias, c in sorted(CITY_ALIASES.items(), key=lambda kv: -len(kv[0])):
            if re.search(rf"\b{re.escape(alias)}\b", low):
                city = c
                break
        if city is None and parts:
            city = parts[0]
    if state is None:
        others = [p for p in parts if canon_city(p) != city and p != city]
        state = canon_state(others[0]) if others else ("Karnataka" if "karnataka" in text.lower() else None)
    return city, state, pin


# ---------------------------------------------------------------- sheet helpers
def _map_headers(row: list[Any], aliases: dict, keywords: list) -> dict[int, str]:
    """Column index -> field key for one candidate header row."""
    lookup = {a: key for key, al in aliases.items() for a in al}
    mapping: dict[int, str] = {}
    for idx, cell in enumerate(row):
        h = norm_header(cell)
        if not h:
            continue
        key = lookup.get(h)
        if key is None:
            close = difflib.get_close_matches(h, list(lookup), n=1, cutoff=0.85)
            key = lookup[close[0]] if close else None
        if key is None:
            for k, words in keywords:
                if all(w in h for w in words):
                    key = k
                    break
        if key and key not in mapping.values():
            mapping[idx] = key
    return mapping


def _find_header(rows: list[list[Any]], aliases: dict, keywords: list) -> tuple[int, dict[int, str]]:
    best = (-1, {})
    for i, row in enumerate(rows[:15]):
        m = _map_headers(row, aliases, keywords)
        if len(m) > len(best[1]):
            best = (i, m)
    return best


def _sheet_rows(ws) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _row_is_example_style(ws, excel_row: int) -> bool:
    try:
        c = ws.cell(excel_row, 1)
        return bool(c.font and c.font.i) or (c.fill is not None and c.fill.fgColor.rgb == "FFF2F2F2")
    except Exception:
        return False


def _classify_sheets(wb) -> dict[str, Any]:
    found: dict[str, Any] = {}
    candidates = []
    for ws in wb.worksheets:
        n = ws.title.lower()
        if any(w in n for w in ("glossary", "guide", "allowed", "readme", "instruction")):
            continue
        if "applic" in n and "applications" not in found:
            found["applications"] = ws
        elif "territor" in n and "territories" not in found:
            found["territories"] = ws
        elif "store" in n and "stores" not in found:
            found["stores"] = ws
        else:
            candidates.append(ws)
    for ws in candidates:  # unnamed sheets: decide by header content
        rows = _sheet_rows(ws)
        scores = {
            "applications": len(_find_header(rows, APP_HEADERS, APP_KEYWORDS)[1]),
            "territories": len(_find_header(rows, TER_HEADERS, TER_KEYWORDS)[1]),
            "stores": len(_find_header(rows, STORE_HEADERS, STORE_KEYWORDS)[1]),
        }
        for kind, score in sorted(scores.items(), key=lambda kv: -kv[1]):
            if score >= 3 and kind not in found:
                found[kind] = ws
                break
    return found


# ---------------------------------------------------------------- applications
def _parse_application(cells: dict[str, Any], excel_row: int) -> dict:
    issues: list[dict] = []
    app: dict[str, Any] = {"row": excel_row, "raw": {k: clean_text(v) for k, v in cells.items()}}
    app["app_id"] = clean_text(cells.get("app_id"))
    app["name"] = clean_text(cells.get("name"))

    city = cells.get("city")
    state = cells.get("state")
    pin = cells.get("pin")
    if "city_state_pin" in cells:
        c2, s2, p2 = split_location(clean_text(cells["city_state_pin"]))
        city, state, pin = city if not is_missing(city) else c2, state if not is_missing(state) else s2, pin if not is_missing(pin) else p2
    if "city_state" in cells:
        c2, s2, _ = split_location(clean_text(cells["city_state"]))
        city = city if not is_missing(city) else c2
        state = state if not is_missing(state) else s2
    app["city"] = canon_city(city)
    app["state"] = canon_state(state)
    app["pin"] = parse_pin(pin)

    inv, inv_status = parse_number(cells.get("investment"))
    if inv_status == "ok" and inv is not None and inv >= 1000:  # looks like rupees, not lakh
        issues.append({"code": "INVESTMENT_RUPEES", "raw": app["raw"].get("investment"), "lakh": round(inv / 100000, 2)})
        inv = inv / 100000
    app["investment"], app["investment_status"] = inv, inv_status

    exp, exp_status = parse_number(cells.get("experience"))
    app["experience"], app["experience_status"] = exp, exp_status
    size, size_status = parse_number(cells.get("size"))
    app["size"], app["size_status"] = size, size_status
    launch, launch_status = parse_number(cells.get("launch"))
    app["launch"], app["launch_status"] = launch, launch_status

    app["business"] = clean_text(cells.get("business")) or ""
    app["notes"] = clean_text(cells.get("notes")) or ""
    app["space"] = canon_space(cells.get("space"))
    app["floor"] = canon_floor(cells.get("floor"))
    app["language"] = canon_language(cells.get("language"))
    app["parse_issues"] = issues
    app["withdrawn"] = False
    app["funds_confirmed"] = False
    app["space_confirmed"] = False
    return app


def _matches_example(app: dict) -> bool:
    ex = TEMPLATE_EXAMPLES["applications"]
    return all((app.get(k) or "") == v for k, v in ex.items())


def parse_applications(ws, warnings: list) -> list[dict]:
    rows = _sheet_rows(ws)
    h_idx, mapping = _find_header(rows, APP_HEADERS, APP_KEYWORDS)
    if h_idx < 0 or len(mapping) < 3:
        warnings.append({"sheet": ws.title, "row": None, "message": "Could not find the Applications header row."})
        return []
    missing_cols = [k for k in ("app_id", "investment", "space") if k not in mapping.values()]
    if not any(k in mapping.values() for k in ("pin", "city_state_pin")):
        missing_cols.append("pin")
    if missing_cols:
        warnings.append({"sheet": ws.title, "row": h_idx + 1,
                         "message": f"Columns not found: {', '.join(missing_cols)}. Those fields are treated as blank."})
    apps: list[dict] = []
    seen_ids: dict[str, int] = {}
    first_data = True
    for i in range(h_idx + 1, len(rows)):
        row = rows[i]
        excel_row = i + 1
        if all(is_missing(v) for v in row):
            continue
        cells = {key: (row[idx] if idx < len(row) else None) for idx, key in mapping.items()}
        try:
            app = _parse_application(cells, excel_row)
        except Exception as e:  # never crash on a bad row
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"Row skipped, could not be read ({e})."})
            continue
        if first_data and _matches_example(app) and _row_is_example_style(ws, excel_row):
            warnings.append({"sheet": ws.title, "row": excel_row,
                             "message": "Template example row (APP-001, Ravi Kumar, grey) skipped; it is not data."})
            first_data = False
            continue
        first_data = False
        if not app["app_id"]:
            app["app_id"] = f"ROW-{excel_row}"
            app["missing_id"] = True
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"No Application ID; labelled {app['app_id']}."})
        if app["app_id"] in seen_ids:
            new_id = f"{app['app_id']} (row {excel_row})"
            warnings.append({"sheet": ws.title, "row": excel_row,
                             "message": f"Application ID {app['app_id']} is used twice; this row is labelled '{new_id}'."})
            app["app_id"] = new_id
        seen_ids[app["app_id"]] = excel_row
        for issue in app["parse_issues"]:
            if issue["code"] == "INVESTMENT_RUPEES":
                warnings.append({"sheet": ws.title, "row": excel_row,
                                 "message": f"{app['app_id']}: investment '{issue['raw']}' looks like rupees; read as ₹{issue['lakh']} lakh."})
        apps.append(app)
    return apps


# ---------------------------------------------------------------- territories
def parse_pin_list(v: Any) -> tuple[list[str], list[str]]:
    if is_missing(v):
        return [], []
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return [str(int(v))], []
    good, bad = [], []
    for tok in re.split(r"[,;/\s\n]+", str(v)):
        tok = parse_pin(tok)
        if not tok:
            continue
        (good if re.fullmatch(r"\d{6}", tok) else bad).append(tok)
    return list(dict.fromkeys(good)), bad


def parse_territories(ws, warnings: list) -> list[dict]:
    cfg = refdata.config()
    rows = _sheet_rows(ws)
    h_idx, mapping = _find_header(rows, TER_HEADERS, TER_KEYWORDS)
    if h_idx < 0 or len(mapping) < 3:
        warnings.append({"sheet": ws.title, "row": None, "message": "Could not find the Territories header row."})
        return []
    out: list[dict] = []
    first_data = True
    seen: set[str] = set()
    for i in range(h_idx + 1, len(rows)):
        row = rows[i]
        excel_row = i + 1
        if all(is_missing(v) for v in row):
            continue
        cells = {key: (row[idx] if idx < len(row) else None) for idx, key in mapping.items()}
        try:
            tid = clean_text(cells.get("territory_id")) or f"TER-ROW{excel_row}"
            city, state = cells.get("city"), cells.get("state")
            if "city_state" in cells:
                c2, s2, _ = split_location(clean_text(cells["city_state"]))
                city = city if not is_missing(city) else c2
                state = state if not is_missing(state) else s2
            pins_text = clean_text(cells.get("pins")) or ""
            pins, bad = parse_pin_list(cells.get("pins"))
            slots, slot_status = parse_number(cells.get("slots"))
            ter = {
                "territory_id": tid, "city": canon_city(city), "state": canon_state(state), "pins": pins,
                "tier": canon_tier(cells.get("tier")), "slots": int(slots) if slot_status == "ok" and slots and slots > 0 else 0,
                "priority": canon_priority(cells.get("priority")), "row": excel_row,
            }
        except Exception as e:
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"Row skipped, could not be read ({e})."})
            continue
        ex = TEMPLATE_EXAMPLES["territories"]
        if first_data and tid == ex["territory_id"] and pins_text == ex["pins_text"] and _row_is_example_style(ws, excel_row):
            warnings.append({"sheet": ws.title, "row": excel_row,
                             "message": "Template example row (TER-01, grey) skipped; it is not data."})
            first_data = False
            continue
        first_data = False
        if bad:
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"{tid}: ignored invalid pin(s) {', '.join(bad)}."})
        if not pins:
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"{tid}: no valid pin codes listed."})
        if slot_status != "ok" or not slots or slots <= 0:
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"{tid}: open slots missing or 0; no approvals possible here."})
        if ter["tier"] not in cfg["cost_model"]:
            ref = refdata.cities().get(ter["city"] or "")
            if ref:
                warnings.append({"sheet": ws.title, "row": excel_row,
                                 "message": f"{tid}: city tier missing/invalid; using {ref['tier']} for {ter['city']}."})
                ter["tier"] = ref["tier"]
        if ter["priority"] not in cfg["priority_points"]:
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"{tid}: priority missing/invalid; treated as Low."})
            ter["priority"] = "Low"
        if tid in seen:
            ter["territory_id"] = f"{tid} (row {excel_row})"
            warnings.append({"sheet": ws.title, "row": excel_row, "message": f"Territory ID {tid} is used twice; this row is labelled '{ter['territory_id']}'."})
        seen.add(ter["territory_id"])
        out.append(ter)
    return out


# ---------------------------------------------------------------- stores
def parse_stores(ws, warnings: list) -> list[dict]:
    rows = _sheet_rows(ws)
    h_idx, mapping = _find_header(rows, STORE_HEADERS, STORE_KEYWORDS)
    if h_idx < 0 or "pin" not in mapping.values():
        return []
    out = []
    for i in range(h_idx + 1, len(rows)):
        row = rows[i]
        if all(is_missing(v) for v in row):
            continue
        cells = {key: (row[idx] if idx < len(row) else None) for idx, key in mapping.items()}
        pin = parse_pin(cells.get("pin"))
        if not pin or not re.fullmatch(r"\d{6}", pin):
            continue
        opened, _ = parse_number(cells.get("opened"))
        out.append({"store_id": clean_text(cells.get("store_id")) or f"ST-{pin}",
                    "city": canon_city(cells.get("city")), "pin": pin, "opened": int(opened) if opened else None})
    return out


# ---------------------------------------------------------------- entry point
def read_workbook(source: Any) -> dict:
    """Read an uploaded workbook (path, bytes or file-like). Returns
    {applications, territories, stores, stores_source, warnings, summary}."""
    warnings: list[dict] = []
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    sheets = _classify_sheets(wb)

    apps = parse_applications(sheets["applications"], warnings) if "applications" in sheets else []
    if "applications" not in sheets:
        warnings.append({"sheet": None, "row": None, "message": "No Applications sheet found."})
    ters = parse_territories(sheets["territories"], warnings) if "territories" in sheets else []
    if "territories" not in sheets:
        warnings.append({"sheet": None, "row": None, "message": "No Territories sheet found."})

    stores = parse_stores(sheets["stores"], warnings) if "stores" in sheets else []
    stores_source = "upload"
    if not stores:
        stores, stores_source = [dict(s) for s in refdata.bundled_stores()], "bundled"
    else:
        bundled = {s["pin"] for s in refdata.bundled_stores()}
        if {s["pin"] for s in stores} != bundled:
            warnings.append({"sheet": "Existing stores", "row": None,
                             "message": "Existing-stores sheet differs from the fixed list; the uploaded sheet is used."})

    rows_with_warnings = len({w["row"] for w in warnings if w.get("row")})
    summary = {"applications": len(apps), "territories": len(ters), "stores": len(stores),
               "stores_source": stores_source, "rows_with_warnings": rows_with_warnings}
    return {"applications": apps, "territories": ters, "stores": stores, "stores_source": stores_source,
            "warnings": warnings, "summary": summary}
