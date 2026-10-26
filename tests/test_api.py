import io
import os
import re
import sys
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.pop("ANTHROPIC_API_KEY", None)

from app import app  # noqa: E402

client = TestClient(app)


def _sample():
    r = client.post("/api/sample")
    assert r.status_code == 200
    return r.json()


def test_index_and_meta():
    r = client.get("/")
    assert r.status_code == 200 and "Partnerly" in r.text and "Franchise Onboarding Engine" in r.text
    m = client.get("/api/meta").json()
    assert m["llm_available"] is False and len(m["worked_examples"]) == 4


def test_sample_counts_and_session_shape():
    j = _sample()
    assert j["result"]["counts"] == {"Approved": 9, "Waitlisted": 7, "On hold": 7, "Not approved": 4}
    assert j["result"]["invariant_errors"] == []
    s = j["session"]
    assert len(s["applications"]) == 27 and len(s["territories"]) == 6 and len(s["signals"]) == 27


def test_upload_file():
    raw = (ROOT / "tests" / "test_data.xlsx").read_bytes()
    r = client.post("/api/upload", files={"file": ("test.xlsx", raw, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["summary"]["applications"] == 27


def test_bad_upload_is_400():
    r = client.post("/api/upload", files={"file": ("x.xlsx", b"not excel", "application/octet-stream")})
    assert r.status_code == 400


def test_run_roundtrip_is_deterministic():
    j = _sample()
    r = client.post("/api/run", json=j["session"]).json()
    assert r["result"] == j["result"]


def test_chat_update_and_knock_on():
    j = _sample()
    r = client.post("/api/chat", json={**j["session"], "message": "APP-012 has now found a ground-floor space"}).json()
    assert r["kind"] == "update"
    recs = r["result"]["records"]
    assert recs["APP-012"]["decision"] == "Approved" and recs["APP-019"]["decision"] == "Waitlisted"
    # the browser sends the updated session back; the update persists
    s2 = {**j["session"], "applications": r["applications"], "signals": r["signals"]}
    q = client.post("/api/chat", json={**s2, "message": "Why is APP-019 waitlisted?"}).json()
    assert q["kind"] == "question" and "APP-012" in q["reply"]
    # the browser also sends its change log, so the assistant can say what changed
    log = [{"time": "10:00", "message": "APP-012 has now found a ground-floor space", "affected": ["APP-012", "APP-019", "APP-020"],
            "per_app": {c["app_id"]: {"old": c["old"], "new": c["new"], "text": c["summary"]} for c in r["changes"]}}]
    h = client.post("/api/chat", json={**s2, "message": "What changed for APP-019 today?", "change_log": log}).json()
    assert h["kind"] == "question" and "Approved → Waitlisted" in h["reply"]
    x = client.post("/api/export", json={**s2, "change_log": log})
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["Change log"]
    assert ws.cell(1, 5).value == "Effect on each applicant" and "APP-020: still Approved; rank" in ws.cell(2, 5).value


def test_chat_question_and_unknown_id():
    j = _sample()
    r = client.post("/api/chat", json={**j["session"], "message": "Why is APP-007 on hold?"}).json()
    assert r["kind"] == "question" and "result" not in r
    r = client.post("/api/chat", json={**j["session"], "message": "APP-099 found space"}).json()
    assert r["reply"] == "There is no APP-099 in the uploaded file."


def test_export_and_template_are_xlsx():
    j = _sample()
    r = client.post("/api/export", json={**j["session"], "change_log": []})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    t = client.get("/api/template")
    assert t.status_code == 200 and t.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Follow-up emails" in wb.sheetnames
    rows = list(wb["Follow-up emails"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == j["result"]["counts"]["On hold"] and all(row[4].startswith("Your franchise application") for row in rows)


def test_sample_file_download_has_expected_results_and_uploads():
    r = client.get("/api/sample-file")
    assert r.status_code == 200 and "Franchise_Sample_Test_Data.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Expected results" in wb.sheetnames
    expected = {row[0]: row[3] for row in wb["Expected results"].iter_rows(min_row=4, values_only=True)
                if row[0] and re.fullmatch(r"APP-\d+", str(row[0]))}
    up = client.post("/api/upload", files={"file": ("s.xlsx", r.content, "application/octet-stream")}).json()
    assert len(expected) == 27
    assert {k: v["decision"] for k, v in up["result"]["records"].items()} == expected
